"""
HealthConnect Rwanda - Week 6
Builds the five Word reports from outputs/results.json, tables and figures

Run from the repository root:  python src/07_build_reports.py
"""
import json, os, glob
import pandas as pd
from docx import Document
from docx.shared import Pt, Inches, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

os.chdir(os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', 'outputs')); os.makedirs('reports', exist_ok=True)
J = json.load(open('results.json')); R, FIGS = J['R'], J['FIGS']
T = lambda n: pd.read_csv(f'tables/{n}.csv')
FK = {k.split('_')[0]: k for k in FIGS}           # f01 -> f01_patient_profile
BLUE = RGBColor(0x2a, 0x78, 0xd6)

def new_doc(title, subtitle):
    d = Document()
    s = d.sections[0]; s.left_margin = s.right_margin = Cm(2); s.top_margin = s.bottom_margin = Cm(2)
    st = d.styles['Normal']; st.font.name = 'Calibri'; st.font.size = Pt(10.5)
    for h, sz in (('Heading 1', 16), ('Heading 2', 13), ('Heading 3', 11)):
        d.styles[h].font.name = 'Calibri'; d.styles[h].font.size = Pt(sz); d.styles[h].font.color.rgb = BLUE
    p = d.add_paragraph(); r = p.add_run(title); r.bold = True; r.font.size = Pt(22); r.font.color.rgb = BLUE
    p = d.add_paragraph(); r = p.add_run(subtitle); r.font.size = Pt(11); r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    p = d.add_paragraph(); r = p.add_run('HealthConnect Rwanda · AnalystLab Africa Experience Lab, Week 6 · Track 2: Healthcare Predictive Analytics and Population Health Intelligence'); r.font.size = Pt(9); r.italic = True
    return d

def para(d, text, bold=False, italic=False, size=None):
    p = d.add_paragraph(); r = p.add_run(text); r.bold = bold; r.italic = italic
    if size: r.font.size = Pt(size)
    p.paragraph_format.space_after = Pt(6); return p

def bullets(d, items):
    for t in items:
        p = d.add_paragraph(style='List Bullet')
        if isinstance(t, tuple):
            r = p.add_run(t[0]); r.bold = True; p.add_run(t[1])
        else: p.add_run(t)
        p.paragraph_format.space_after = Pt(2)

def shade(cell, hexcol):
    tcPr = cell._tc.get_or_add_tcPr(); sh = OxmlElement('w:shd')
    sh.set(qn('w:val'), 'clear'); sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), hexcol); tcPr.append(sh)

def table(d, df, widths=None, size=8.5, maxrows=None):
    if maxrows: df = df.head(maxrows)
    t = d.add_table(rows=1, cols=len(df.columns)); t.style = 'Table Grid'; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, c in enumerate(df.columns):
        cell = t.rows[0].cells[i]; cell.text = ''; r = cell.paragraphs[0].add_run(str(c)); r.bold = True; r.font.size = Pt(size); r.font.color.rgb = RGBColor(255, 255, 255); shade(cell, '2A78D6')
    for _, row in df.iterrows():
        cells = t.add_row().cells
        for i, v in enumerate(row):
            if isinstance(v, float):
                v = '' if pd.isna(v) else (f'{v:,.0f}' if abs(v) >= 1000 else f'{v:.3g}' if abs(v) < 1 else f'{v:,.2f}')
            cells[i].text = ''; r = cells[i].paragraphs[0].add_run('' if v is None or (isinstance(v, float) and pd.isna(v)) else str(v)); r.font.size = Pt(size)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths): row.cells[i].width = Cm(w)
    d.add_paragraph().paragraph_format.space_after = Pt(2)
    return t

def fig(d, key, width=17):
    name = FK[key]; d.add_heading(f"Figure {key[1:]}. {FIGS[name]['title']}", 3)
    d.add_picture(f'figures/{name}.png', width=Cm(width)); d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

def finish(d, fn):
    d.save(f'reports/{fn}'); print('saved', fn)

bn = lambda x: f'{x/1e9:.2f} bn'
sev = R['dq_counts']

# ------------------------------------------------------------------ 1. DATA QUALITY
d = new_doc('Data Quality Report', 'Loading, merging, cleaning and validation of the HealthConnect Rwanda data')
d.add_heading('1. Scope and approach', 1)
para(d, f"{R['n_source_files']} source files ({R['n_tables_loaded']} tables/sheets) were loaded and profiled: patient demographics and geography, appointments, diagnoses, prescriptions, lab tests, admissions, referrals, billing, payments, insurance claims, feedback, staff, departments, rooms, equipment, facilities and revenue summaries. "
          f"Every table was checked for completeness, duplicates, validity, range plausibility and referential integrity. In total {R['dq_n_checks']} checks were run: {sev.get('Critical',0)} Critical, {sev.get('High',0)} High, {sev.get('Medium',0)} Medium, {sev.get('Low',0)} Low and {sev.get('Info',0)} Informational (checks that passed or are descriptive).")
d.add_heading('2. Headline finding', 1)
para(d, "The data are internally clean at row level (few missing values, no duplicate keys) but are not relationally consistent. The main problems are structural and cannot be repaired by cleaning; they limit what any analysis can conclude and are documented so that results are interpreted correctly.")
bullets(d, [("Only 5,000 of 60,000 billed visits have an appointment record. ", "The Appointments sheet holds 5,000 rows, while billing, payments, claims and diagnoses hold 60,000. Analyses needing clinical detail use the 5,000-appointment cohort; financial analyses use all 60,000 bills."),
            ("Facility, department and doctor keys disagree. ", "About 95% of appointments reference a department or doctor that belongs to a different facility, so IDs appear to have been generated independently. Facility is taken from the appointment record; department type from the department ID (a documented assumption)."),
            ("Bill 'Status' is unreliable. ", "80% of bills are paid in full whatever the Status says; 'Pending' bills have the same average payment ratio as 'Paid'."),
            ("Revenue_Summary does not reconcile with billing. ", f"Monthly summary total {bn(R['rev_total'])} RWF versus {bn(R['bill_total'])} RWF billed."),
            ("No facility master, bed status or equipment availability fields. ", "All 2,000 beds are flagged 'Available', so real bed use cannot be assessed.")])
d.add_heading('3. Referential-integrity checks', 1)
rel = pd.DataFrame(R['relationships']); rel['Match_Rate'] = (rel['Match_Rate'] * 100).round(1).astype(str) + '%'
table(d, rel, [8, 2.5, 2.5, 2.5])
d.add_heading('4. Data quality findings by severity', 1)
dq = T('data_quality_findings')
dq['Pct_Affected'] = dq['Pct_Affected'].round(1)
cols = ['Table', 'Check', 'Records_Checked', 'Issues_Found', 'Pct_Affected', 'Action_Taken']
for s_, h in (('Critical', 'Critical'), ('High', 'High'), ('Medium', 'Medium'), ('Low', 'Low')):
    sub = dq[dq.Severity == s_]
    if len(sub):
        d.add_heading(f'4.{["Critical","High","Medium","Low"].index(s_)+1} {h} ({len(sub)})', 2)
        table(d, sub[cols].rename(columns={'Records_Checked': 'Checked', 'Issues_Found': 'Issues', 'Pct_Affected': '% affected', 'Action_Taken': 'Action taken'}), [2.6, 5, 1.6, 1.6, 1.4, 5.2], size=8)
para(d, f"The remaining {sev.get('Info',0)} informational checks (passed checks and descriptive counts) are listed in outputs/tables/data_quality_findings.csv.", italic=True)
d.add_heading('5. Cleaning decisions', 1)
bullets(d, ["Missing values: 'None' in ChronicCondition is a valid category (no chronic disease), not a missing value, and was preserved. No other column had material missingness.",
            "Implausible values were flagged, not deleted: anthropometrics (10.8% of patients), extreme BMI (5.8%) and age/date-of-birth inconsistencies (18.4%). Flags are carried as columns so analyses can include or exclude them.",
            "Orphan child records (diagnoses, payments, claims without a parent appointment) were retained for financial analysis and excluded from joined clinical analyses; the n for each analysis is disclosed.",
            "Data types were standardised (dates, categorical codes, numeric amounts); text categories were trimmed and case-normalised.",
            "Privacy: names, phone numbers and emergency contacts are excluded from all exported datasets."])
d.add_heading('6. Integrated datasets produced', 1)
table(d, pd.DataFrame([
    ['HC_Integrated_Encounters_5000.csv', f"{R['n_encounters']:,} appointments joined to patient, geography, doctor, department, diagnosis, prescription, lab, admission, referral, bill and feedback; engineered features"],
    ['HC_Integrated_Billing_Feedback_60000.csv', f"{R['n_fin']:,} bills joined to payments, insurance claims, diagnoses and feedback"],
    ['HC_Clean_Patients_20000.csv', f"{R['n_patients']:,} patients with geography, risk flags and engineered features (no direct identifiers)"],
    ['HC_Clean_Staff_1250.csv', '1,250 staff with department and facility']], columns=['File', 'Content']), [6, 11])
d.add_heading('7. Impact on analysis and recommendations', 1)
para(d, "Because outcome variables (admission, satisfaction, bill amount, length of stay) show no relationship with any recorded driver, and because key links between tables are broken, conclusions in the other reports are deliberately conservative. The first recommendation to HealthConnect is to fix the data foundation: re-extract the full-year Appointments table, enforce Facility→Department→Doctor foreign keys, derive bill Status from amounts, and add facility master, bed-status and equipment-condition fields.")
finish(d, '1_Data_Quality_Report.docx')

# ------------------------------------------------------------------ 2. EDA
d = new_doc('Exploratory Data Analysis Report', 'Patient, clinical, operational, financial and patient-experience analysis')
d.add_heading('1. Summary', 1)
bullets(d, [f"{R['n_patients']:,} patients; mean age {R['mean_age']:.1f}; {R['gender_split']['Male']:,} male and {R['gender_split']['Female']:,} female; {R['chronic_prev']:.1%} have a chronic condition; {R['highrisk_pct']:.1%} meet the high-risk rule.",
            f"Patients are spread evenly over the five provinces (Western {R['province_patients']['Western']:,}; Northern {R['province_patients']['Northern']:,}) and every patient has insurance.",
            f"Clinical: five diagnoses are almost equally common (about 12,000 each) and severity is split evenly across Mild/Moderate/Severe. Admission rate {R['admit_rate']:.1%} (586 of 5,000), mean length of stay {R['mean_los']:.1f} days, referral rate {R['referral_rate']:.1%}.",
            f"Financial: {bn(R['bill_total'])} RWF billed, {bn(R['paid_total'])} paid, {bn(R['outstanding_total'])} outstanding ({R['outstanding_pct']:.1f}%). Claim approval is {R['claim_approval']:.0%}; payment methods are evenly split.",
            f"Experience: mean rating {R['mean_rating']:.2f}/5 with ratings spread evenly from 1 to 5; 40% rate 1–2. Feedback exists for {R['feedback_coverage']:.0%} of bills.",
            "A recurring theme: categories are almost uniformly distributed, so the data behave as if generated at random. Differences between groups, facilities and departments are small and mostly within chance variation."])
sections = [('2. Patient analysis (demographics, geography, insurance, growth)', ['f01', 'f02', 'f03', 'f04', 'f05', 'f06']),
            ('3. Clinical analysis (diagnoses, labs, prescriptions, admissions, referrals)', ['f07', 'f08', 'f09']),
            ('4. Operational analysis (staff workload, facility utilisation, capacity)', ['f10', 'f11', 'f12']),
            ('5. Financial analysis (revenue, billing, claims, payments)', ['f13', 'f14', 'f15']),
            ('6. Patient experience analysis (satisfaction)', ['f16', 'f17'])]
for h, ks in sections:
    d.add_heading(h, 1)
    for k in ks:
        fig(d, k)
d.add_heading('7. Engineered features', 1)
para(d, 'Ten features were created and saved in the integrated datasets:')
bullets(d, ['Age group (7 bands) and BMI category', 'Chronic-disease flag and number of conditions', 'Clinical risk score and high-risk flag (65+/50–64 with chronic disease, obesity or underweight)',
            'Visit frequency and recency (days since last appointment)', 'Length of stay and admission flag', 'Bill amount per visit, payment ratio and outstanding balance',
            'Claim coverage ratio (claim amount ÷ bill)', 'Doctor workload (appointments per clinician) and department workload index', 'Facility performance score (0–100: 35% satisfaction, 30% collection rate, 15% claim approval, 20% load per clinician)', 'Satisfaction group (poor 1–2, neutral 3, good 4–5)'])
finish(d, '2_EDA_Report.docx')

# ------------------------------------------------------------------ 3. STATISTICS
d = new_doc('Statistical Analysis Report', 'Descriptive statistics, correlations and hypothesis tests')
d.add_heading('1. Summary', 1)
bullets(d, [f"{R['ht_total']} hypothesis tests were run at α = 0.05; {R['ht_rejected']} rejected the null hypothesis, also after Holm correction for multiple testing.",
            f"Of 26 candidate admission factors screened, {R['admission_factors_sig_holm']} remained significant after correction.",
            "Eleven correlation pairs were examined. The only strong, significant association is facility billed revenue ↔ appointment count (r = 0.90), which is arithmetic: more visits means more bills.",
            "Interpretation: apart from mechanical links (BMI from height and weight, revenue from volume), recorded variables are statistically independent. This is a data finding as much as a statistical one."])
d.add_heading('2. Descriptive statistics', 1)
ds = T('descriptive_statistics'); table(d, ds, size=7)
d.add_heading('3. Distributions', 1)
for k in ['f18', 'f19']: fig(d, k)
d.add_heading('4. Correlation analysis', 1)
c = T('correlations'); c['p_value'] = c['p_value'].map(lambda x: f'{x:.3g}'); c['Coefficient'] = c['Coefficient'].round(3)
table(d, c[['Pair', 'Data', 'n', 'Method', 'Coefficient', 'p_value', 'Strength', 'Significant (α=.05)']].rename(columns={'Significant (α=.05)': 'Sig.'}), [4.2, 2.8, 1.2, 3.5, 1.6, 1.4, 1.6, 1])
for k in ['f20', 'f21', 'f22']: fig(d, k)
d.add_heading('5. Admission and chronic-disease risk factors', 1)
for k in ['f23', 'f24', 'f25']: fig(d, k)
d.add_heading('6. Hypothesis tests', 1)
para(d, 'Each test states H0 and H1, the test used, α, sample size, test statistic, p-value, effect size and decision. Holm-adjusted p-values control the family-wise error rate across all tests.')
h = T('hypothesis_tests')
for _, r in h.iterrows():
    d.add_heading(f"{r['ID']}. {r['Question']}", 3)
    table(d, pd.DataFrame([['H0', r['H0']], ['H1', r['H1']], ['Test', r['Test']], ['α / n', f"{r['Alpha']} / {int(r['n']):,}"], ['Statistic', f"{r['Statistic']:.3g}"],
                           ['p-value (Holm)', f"{r['p_value']:.3g} ({r['p_holm_all']:.3g})"], ['Effect size', f"{r['Effect']:.3g} ({r['Effect_measure']})"], ['Decision', r['Decision']], ['Conclusion', r['Interpretation']]], columns=['Item', 'Detail']), [3.2, 13.8], size=8.5)
fig(d, 'f26')
finish(d, '3_Statistical_Analysis_Report.docx')

# ------------------------------------------------------------------ 4. PREDICTIVE
d = new_doc('Predictive Analytics Report', 'Forecasting, classification, regression, clustering and PCA')
d.add_heading('1. Summary', 1)
fa, fr = R['fc_app'], R['fc_rev']
table(d, pd.DataFrame([
    ['Visit forecast (12 months)', f"Simple exponential smoothing ≈ 4,975 visits/month (≈59,700 in 2027)", f"MAPE {fa['mape_best']:.1f}% vs {fa['mape_mean']:.1f}% for the historical mean", 'Useful only as a flat baseline; no trend'],
    ['Revenue forecast (12 months)', f"≈ {bn(1.10e9)} RWF for 2027 (≈92 M/month)", f"MAPE {fr['mape_best']:.0f}% vs {fr['mape_mean']:.0f}% for the mean", 'High uncertainty (95% band ±68%)'],
    ['Admission prediction', f"Best CV AUC {R['admission_best']['AUC']:.3f}", f"Permutation p = {R['admission_perm_p']:.2f}; hold-out AUC {R['admission_holdout']['auc']:.3f} ({R['admission_holdout']['lo']:.2f}–{R['admission_holdout']['hi']:.2f})", 'No predictive signal (chance = 0.50)'],
    ['Satisfaction prediction', f"Best real-model R² {R['sat_best']['R2']:.3f}", f"Permutation p = {R['sat_perm_p']:.2f}; hold-out R² {R['sat_holdout_r2']:.3f}", 'No predictive signal'],
    ['Department workload forecast', 'Cardiology and Radiology highest (load index 109 and 105)', 'Monte-Carlo overload probability 26%–38%', 'Differences not significant'],
    ['Clustering', 'Patients: 3 chronic-condition groups + none; Facilities: 2 volume-based groups; Departments: continuous', 'Silhouette 0.28 / 0.19 / 0.14', 'Weak structure; for peer grouping only']],
    columns=['Task', 'Result', 'Validation', 'Verdict']), [3.2, 5.4, 5.2, 3.2], size=8.5)
para(d, "Method principle: every model is compared with a naive baseline and with a label-permutation null, using cross-validation and (for classification) a held-out test set. A model is only reported as useful if it beats those benchmarks. Statsmodels and Jupyter could not be installed in the analysis sandbox, so exponential-smoothing models were implemented directly in NumPy and the notebook was executed with a custom runner; the notebook file is a standard .ipynb with all outputs saved.", italic=True)
d.add_heading('2. Forecasting (visits, revenue, department workload)', 1)
para(d, "With only 12 monthly observations, models were chosen by rolling-origin back-testing (minimum 6 months training, horizons 1–3). The Mean, Naïve, Linear trend, Simple Exponential Smoothing and damped Holt models were compared; 95% bands come from back-test RMSE.")
bt = pd.DataFrame(fa['backtest'])[['Model', 'MAE', 'RMSE', 'MAPE']]; d.add_heading('Visit forecast back-test', 3); table(d, bt)
bt = pd.DataFrame(fr['backtest'])[['Model', 'MAE', 'RMSE', 'MAPE']]; d.add_heading('Revenue forecast back-test', 3); table(d, bt)
for k in ['f27', 'f28', 'f29']: fig(d, k)
d.add_heading('Department workload forecast', 3); table(d, pd.DataFrame(R['dept_forecast']).round(2), size=8)
d.add_heading('3. Hospital admission prediction (classification)', 1)
para(d, 'Three feature sets were tested: A booking-time information only; B adding diagnosis and severity; C adding labs, prescriptions and referral. Models: logistic regression, random forest, gradient boosting, against a prior-probability baseline (repeated stratified CV).')
am = T('admission_model_cv'); table(d, am.round(3), size=8)
for k in ['f30', 'f31']: fig(d, k)
d.add_heading('4. Patient satisfaction prediction', 1)
sm = T('satisfaction_model_cv'); table(d, sm.round(3), size=7.5)
fig(d, 'f32')
d.add_heading('5. Clustering and segmentation', 1)
pc = R['patient_clusters']; para(d, f"Patients: silhouette peaks at {pc['sil_best']:.2f} (k = {pc['k_best']}) against {pc['null_sil_at_best']:.2f} for shuffled data. The 4-segment solution simply reproduces the chronic-condition indicators. High-risk shares are 38–39% in each condition group and 18% in the no-chronic-disease group.")
d.add_heading('Patient segments (k = 4)', 3); table(d, pd.DataFrame(pc['profile'])[['Label', 'Patients', 'MeanAge', 'MeanBMI', 'Female', 'MeanVisits', 'HighRisk', 'AdmitRate', 'Description']].round(3), size=8)
fc = R['facility_clusters']; d.add_heading('Facility clusters', 3); table(d, pd.DataFrame(fc['profile'])[['Cluster', 'Facilities', 'Appointments', 'ApptPerClinician', 'StaffPer1000Pts', 'MeanRating', 'CollectionRate', 'PerformanceScore', 'Members']].round(2), size=8)
for k in ['f33', 'f34', 'f35']: fig(d, k)
d.add_heading('6. Principal Component Analysis (optional)', 1)
fig(d, 'f36')
d.add_heading('7. Limitations and what would make the models work', 1)
bullets(d, ["Outcomes are statistically independent of all recorded drivers, so no algorithm can predict them; the models were run to prove this rigorously, not to force a result.",
            "Only 12 months of aggregate history: seasonality and trend cannot be assessed; forecasts are flat baselines with wide bands.",
            "Only 5,000 of 60,000 visits have an appointment record, and facility/department/doctor links are inconsistent.",
            "To enable genuine prediction: collect longitudinal patient histories, wait times, readmission and discharge outcomes, richer feedback items, and 24+ months of data."])
finish(d, '4_Predictive_Analytics_Report.docx')

# ------------------------------------------------------------------ 5. EXECUTIVE SUMMARY
d = new_doc('Executive Summary', 'Healthcare analytics findings and recommendations for HealthConnect Rwanda leadership')
d.add_heading('Purpose and scope', 1)
para(d, f"HealthConnect Rwanda asked whether its operational data can support predictive analytics and population health intelligence. We integrated {R['n_tables_loaded']} tables covering {R['n_patients']:,} patients, 5,000 appointment encounters and 60,000 bills across 20 facilities, ran {R['dq_n_checks']} data-quality checks, built 36 visualisations, ran {R['ht_total']} hypothesis tests and fitted forecasting, classification, regression, clustering and PCA models. All models were benchmarked against naive baselines and permutation tests so that only real signal is reported.")
d.add_heading('Key findings', 1)
bullets(d, [("Data foundation is the biggest constraint. ", f"Only 5,000 of 60,000 billed visits have an appointment record, 95% of appointments point to a department or doctor from another facility, and bill 'Status' does not reflect payment. Revenue_Summary ({bn(R['rev_total'])} RWF) does not reconcile with billing ({bn(R['bill_total'])} RWF)."),
            ("The population is large, chronic and evenly spread. ", f"Mean age {R['mean_age']:.0f}; {R['chronic_prev']:.0%} have a chronic condition and {R['highrisk_pct']:.0%} ({6662:,} patients) meet the clinical high-risk rule, yet only 1,498 of them had an appointment in 2026. Patients, diagnoses and severity are close to uniform across provinces."),
            ("Money is being left uncollected. ", f"{bn(R['outstanding_total'])} RWF ({R['outstanding_pct']:.1f}%) of billed amounts is unpaid, 21,028 insurance claims worth 2.44 bn RWF are pending, and only 50% of claims are approved, with no difference between the four insurers."),
            ("Demand is stable but forecastable only as an average. ", "About 4,975 visits/month (≈59,700 in 2027, typical error ±14%) and ≈1.10 bn RWF revenue (error ±26%). No trend is detectable in 12 months."),
            ("Admissions and satisfaction cannot be predicted from current data. ", f"Best admission model AUC {R['admission_best']['AUC']:.2f} (chance is 0.50; permutation p = {R['admission_perm_p']:.2f}); satisfaction R² ≈ 0. None of {R['ht_total']} hypothesis tests rejected its null hypothesis. This is a finding about the data, not a modelling failure."),
            ("Staffing does not follow volume. ", "Facility volumes differ by 28% (217–278 appointments) while headcount is 61–64 everywhere; staff numbers do not track patient numbers."),
            ("Patient experience needs a better instrument. ", f"Mean rating {R['mean_rating']:.2f}/5; 40% rate 1–2; ratings are unrelated to comments, diagnosis, billing, department or facility. With ≈89 ratings per site, facility rankings are not reliable.")])
d.add_heading('Business recommendations', 1)
rec = T('recommendations')
table(d, rec[['Theme', 'Recommendation', 'Strength of evidence', 'Suggested action']], [2.6, 4.2, 3.4, 6.8], size=8)
d.add_heading('Answers to the strategic questions', 1)
qa = [['Which factors influence hospital admissions?', 'None can be confirmed. 26 factors screened; none significant after correction; best model AUC ≈ 0.51 (chance).'],
      ['Which patient groups are most at risk of chronic disease?', 'Prevalence is ~75% in every group; only blood group differs significantly (negligible effect). Use the clinical risk rule (65+ with chronic disease, obesity or underweight) for outreach.'],
      ['What factors affect patient satisfaction?', 'None identifiable in recorded data (R² ≈ 0); the feedback instrument itself looks unreliable.'],
      ['Can future demand be forecast accurately?', 'Only as a stable average (~5,000 visits/month, ±14%); seasonality cannot be assessed with 12 months.'],
      ['Which departments may become overloaded?', 'None clearly. Cardiology and Radiology have the highest forecast load index (≈109, 105) but differences are not significant.'],
      ['How can resources be allocated more efficiently?', 'Align staffing with facility volume; capture a bed census before changing capacity.'],
      ['What improvements maximise outcomes?', 'Fix data integrity, then collections (0.92 bn unpaid), claims turnaround, a lab-ordering audit and a redesigned feedback survey.']]
table(d, pd.DataFrame(qa, columns=['Question', 'Answer']), [5.5, 11.5], size=8.5)
d.add_heading('Priorities for the next 90 days', 1)
bullets(d, ['Weeks 1–4: re-extract full Appointments data, enforce foreign keys, derive bill Status from amounts, reconcile revenue summary with billing.',
            'Weeks 2–8: weekly outstanding-balance ageing report, follow-up team for pending claims, agreed turnaround times with all four insurers.',
            'Weeks 4–12: launch the new feedback survey (wait time, courtesy, cleanliness, drug availability), outreach to high-risk patients without a 2026 visit, staffing pilot at the busiest and quietest sites.',
            'Then: collect 24+ months of data and re-run the predictive models.'])
d.add_heading('Limitations', 1)
para(d, "Results rest on a data set whose tables do not link consistently, with 12 months of history and no facility master data. 'No effect found' means no effect detectable in these data, not proof of absence. Power BI and SQL deliverables belong to other tracks and are not included.")
finish(d, '5_Executive_Summary.docx')
