"""
HealthConnect Rwanda - Week 6
Builds the PowerPoint deck from outputs/results.json and figures

Run from the repository root:  python src/08_build_slides.py
"""
import json, os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
os.chdir(os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', 'outputs')); os.makedirs('reports', exist_ok=True)
J = json.load(open('results.json')); R, FIGS = J['R'], J['FIGS']
BLUE, DARK, GREY = RGBColor(0x2a,0x78,0xd6), RGBColor(0x1e,0x25,0x2f), RGBColor(0x55,0x5b,0x66)
P = Presentation(); P.slide_width, P.slide_height = Inches(13.333), Inches(7.5)
blank = P.slide_layouts[6]
def tb(s, x, y, w, h, text, size=16, bold=False, color=DARK):
    t = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); f = t.text_frame; f.word_wrap = True
    lines = text if isinstance(text, list) else [text]
    for i, l in enumerate(lines):
        p = f.paragraphs[0] if i == 0 else f.add_paragraph(); p.text = l; p.font.size = Pt(size); p.font.bold = bold; p.font.color.rgb = color; p.space_after = Pt(8)
    return t
def slide(title):
    s = P.slides.add_slide(blank)
    bar = s.shapes.add_shape(1, 0, 0, P.slide_width, Inches(0.12)); bar.fill.solid(); bar.fill.fore_color.rgb = BLUE; bar.line.fill.background()
    tb(s, 0.6, 0.3, 12, 0.9, title, 28, True, BLUE); return s
def bl(items): return ['• ' + i for i in items]
def figslide(title, key, note):
    s = slide(title); name = [k for k in FIGS if k.startswith(key)][0]
    from PIL import Image
    w, h = Image.open(f'figures/{name}.png').size
    mw, mh = 8.4, 5.7; sc = min(mw / w, mh / h)
    s.shapes.add_picture(f'figures/{name}.png', Inches(0.5), Inches(1.3), Inches(w * sc), Inches(h * sc))
    tb(s, 9.2, 1.5, 3.8, 5, bl(note), 15); return s
bn = lambda x: f'{x/1e9:.2f} bn'

s = P.slides.add_slide(blank)
bg = s.shapes.add_shape(1, 0, 0, P.slide_width, P.slide_height); bg.fill.solid(); bg.fill.fore_color.rgb = BLUE; bg.line.fill.background()
tb(s, 0.8, 2.3, 11.5, 1.5, 'HealthConnect Rwanda: Healthcare Predictive Analytics & Population Health Intelligence', 36, True, RGBColor(255,255,255))
tb(s, 0.8, 4.4, 11, 1.5, ['AnalystLab Africa · Experience Lab Week 6 · Data Science Track', 'Findings, models and recommendations'], 20, False, RGBColor(235,242,252))

s = slide('Context and data'); tb(s, 0.6, 1.4, 12, 5.5, bl([
    f"{R['n_source_files']} source files ({R['n_tables_loaded']} tables): {R['n_patients']:,} patients, 5,000 appointments, 60,000 bills, 20 facilities, 1,250 staff",
    'Goal: assess whether the data support demand forecasting, risk prediction and resource planning',
    'Approach: clean and integrate, explore, test hypotheses, model, then recommend',
    'Every model is compared with a naive baseline and a permutation null, so only real signal is reported',
    'Privacy: names, phones and emergency contacts excluded from all exports']), 20)

s = slide('Data quality: structurally inconsistent'); sev = R['dq_counts']
tb(s, 0.6, 1.4, 12, 5.5, bl([f"{R['dq_n_checks']} checks: {sev['Critical']} critical, {sev['High']} high, {sev['Medium']} medium, {sev['Low']} low",
    'Only 5,000 of 60,000 billed visits have an appointment record',
    '95% of appointments reference a department/doctor from another facility',
    "80% of bills are paid in full whatever 'Status' says",
    f"Revenue summary {bn(R['rev_total'])} vs billing {bn(R['bill_total'])} RWF: no reconciliation",
    'Rows are clean (few missing values); links between tables are not']), 20)

figslide('Who are the patients?', 'f01', [f"Mean age {R['mean_age']:.0f}; near 50/50 gender", f"{R['chronic_prev']:.0%} have a chronic condition", f"{R['highrisk_pct']:.0%} meet the high-risk rule", 'Evenly spread across 5 provinces; 100% insured'])
figslide('Revenue and collections', 'f14', [f"Billed {bn(R['bill_total'])} RWF, paid {bn(R['paid_total'])}", f"{bn(R['outstanding_total'])} outstanding ({R['outstanding_pct']:.1f}%)", '50% claim approval; 21,028 claims pending (2.44 bn)', 'No difference between insurers'])
figslide('Patient experience', 'f16', [f"Mean rating {R['mean_rating']:.2f}/5; 40% rate 1–2", 'Ratings spread evenly 1–5', 'Unrelated to comment text, diagnosis, billing, department or facility', 'Instrument needs redesign'])
figslide('Statistical testing', 'f26', [f"{R['ht_total']} hypothesis tests at α = 0.05", f"{R['ht_rejected']} rejected after correction", 'Admission factors: none of 26 significant', 'Only strong correlation is mechanical (revenue vs volume)'])
figslide('Visit forecast, 2027', 'f27', ['≈4,975 visits/month (≈59,700/yr)', f"MAPE {R['fc_app']['mape_best']:.1f}% vs {R['fc_app']['mape_mean']:.1f}% for the mean", 'No trend in 12 months', 'Plan on a flat baseline'])
figslide('Revenue forecast, 2027', 'f28', ['≈1.10 bn RWF (≈92 M/month)', f"MAPE {R['fc_rev']['mape_best']:.0f}%; 95% band ±68%", 'Sep-2026 peak and Dec-2026 trough need explaining', 'Budget flat with a contingency reserve'])
figslide('Admission prediction', 'f30', [f"Best CV AUC {R['admission_best']['AUC']:.2f} (chance 0.50)", f"Permutation p = {R['admission_perm_p']:.2f}", f"Hold-out AUC {R['admission_holdout']['auc']:.2f}", 'Diagnosis, labs, Rx add nothing: no usable signal'])
figslide('Satisfaction prediction', 'f32', [f"Best real-model R² {R['sat_best']['R2']:.3f}", 'Worse than predicting the mean', 'Satisfaction drivers are not recorded', 'Collect wait time, courtesy, cleanliness'])
figslide('Facility clusters and workload', 'f34', ['Two peer groups, split by volume and load', 'Not quality or outcome groups', 'Staff 61–64 per site vs 28% volume spread', 'Cardiology and Radiology: highest forecast load'])

import pandas as pd
rec = pd.read_csv('tables/recommendations.csv')
s = slide('Recommendations')
rows = rec.iloc[[0, 8, 5, 3, 2, 10]]
t = s.shapes.add_table(len(rows) + 1, 2, Inches(0.6), Inches(1.3), Inches(12.1), Inches(5.5)).table
t.columns[0].width = Inches(4.2); t.columns[1].width = Inches(7.9)
for i, h in enumerate(['Recommendation', 'Action']): t.cell(0, i).text = h
for r, (_, x) in enumerate(rows.iterrows(), 1):
    t.cell(r, 0).text = x['Recommendation']; t.cell(r, 1).text = x['Suggested action']
for row in t.rows:
    for c in row.cells:
        for p in c.text_frame.paragraphs: p.font.size = Pt(12)

s = slide('Limitations and next steps'); tb(s, 0.6, 1.4, 12, 5.5, bl([
    'Poor linkage between tables and 12 months of history limit what can be concluded',
    "'No effect found' means none detectable in these data, not proof of absence",
    'Next: fix data integrity, collect 24+ months, redesign feedback, add bed census and wait-time data',
    'Then re-run the notebook: models and reports regenerate automatically',
    'Deliverables: notebook, cleaned datasets, 4 reports, executive summary, 36 figures']), 20)
P.save('reports/6_Presentation_Slides.pptx'); print('slides', len(P.slides))
